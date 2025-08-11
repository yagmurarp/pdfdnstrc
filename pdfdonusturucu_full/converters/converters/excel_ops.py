# converters/excel_ops.py
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
import pandas as pd
import openpyxl
import os

def excel_to_pdf(in_path: str, out_path: str):
    """XLSX -> PDF (sayfaları tablo olarak basar)"""
    xls = pd.ExcelFile(in_path)
    styles = getSampleStyleSheet()
    story = []
    doc = SimpleDocTemplate(out_path, pagesize=A4, rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)

    for i, sheet in enumerate(xls.sheet_names):
        df = xls.parse(sheet)
        story.append(Paragraph(f"Sayfa: {sheet}", styles["Heading3"]))
        data = [df.columns.tolist()] + df.fillna("").astype(str).values.tolist()
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e9edf7')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.black),
            ('GRID', (0,0), (-1,-1), 0.25, colors.grey),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ]))
        story.append(table)
        if i < len(xls.sheet_names) - 1:
            story.append(PageBreak())
    doc.build(story)

def excel_to_word(in_path: str, out_path: str):
    """XLSX -> DOCX (ilk sayfayı tablo yapar)"""
    df = pd.read_excel(in_path, sheet_name=0)
    doc = Document()
    table = doc.add_table(rows=len(df.index)+1, cols=len(df.columns))
    table.style = 'Table Grid'
    # header
    for j, col in enumerate(df.columns):
        table.cell(0, j).text = str(col)
    # rows
    for i, row in enumerate(df.fillna("").astype(str).values, start=1):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    doc.save(out_path)
